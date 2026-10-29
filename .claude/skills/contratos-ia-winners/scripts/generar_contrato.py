#!/usr/bin/env python3
"""
Genera un contrato en Word (.docx) y PDF a partir de un archivo Markdown sencillo.

FORMATO DEL ARCHIVO DE ENTRADA
------------------------------
Empieza con un bloque de metadatos entre líneas '---' y sigue con el cuerpo:

    ---
    titulo: CONTRATO DE PRESTACIÓN DE SERVICIOS DE INTELIGENCIA ARTIFICIAL
    proyecto: Agente de WhatsApp para reservas
    proveedor: Ana López García
    proveedor_firmante: Ana López García
    proveedor_cargo: Titular
    cliente: Clínica Ejemplo S.L.
    cliente_firmante: Juan Pérez
    cliente_cargo: Administrador
    lugar: Madrid
    fecha: 2026-10-29
    revision_juridica: no        # "si" añade el aviso REVISIÓN JURÍDICA RECOMENDADA en la portada
    ---

    # REUNIDOS
    De una parte, **Ana López García**, con NIF ...

    # CLÁUSULAS
    ## 1. Objeto
    1.1. El presente contrato ...

    - elemento de lista
    | Columna | Columna |
    |---|---|
    | celda | celda |

Sintaxis admitida en el cuerpo:
  '# '   título de sección (REUNIDOS, EXPONEN, CLÁUSULAS, ANEXO I...)
  '## '  título de cláusula
  '### ' subtítulo
  '- '   lista con viñetas
  '| '   tabla (la segunda fila '|---|' se ignora)
  '**negrita**' y '*cursiva*' dentro del texto
  Una línea en blanco separa párrafos. El bloque de firmas se añade solo.

USO
---
    python3 generar_contrato.py contrato.md --salida /mnt/user-data/outputs/
    python3 generar_contrato.py contrato.md --salida . --sin-pdf
"""

import argparse
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

COLOR_TITULO = RGBColor(0x1F, 0x2A, 0x44)
COLOR_SUAVE = RGBColor(0x59, 0x5F, 0x6B)
COLOR_AVISO = RGBColor(0xA3, 0x1D, 0x1D)
FUENTE = "Calibri"
AVISO_BORRADOR = ("Borrador informativo generado con ayuda de IA. "
                  "No sustituye la revisión de un profesional jurídico.")
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


# ---------------------------------------------------------------- utilidades

def leer_entrada(ruta):
    with open(ruta, encoding="utf-8") as f:
        texto = f.read()
    meta, cuerpo = {}, texto
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n(.*)$", texto, re.S)
    if m:
        for linea in m.group(1).splitlines():
            linea = linea.split("#", 1)[0].rstrip() if " #" in linea else linea
            if ":" in linea:
                k, v = linea.split(":", 1)
                meta[k.strip().lower()] = v.strip()
        cuerpo = m.group(2)
    return meta, cuerpo


def slug(texto, maximo=30):
    t = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    t = re.sub(r"[^A-Za-z0-9]+", "", t.title())
    return t[:maximo] or "SinDato"


def fecha_larga(iso):
    try:
        a, m, d = (int(x) for x in iso.split("-"))
        return f"{d} de {MESES[m - 1]} de {a}"
    except Exception:
        return iso


def fuente_base(doc):
    estilo = doc.styles["Normal"]
    estilo.font.name = FUENTE
    estilo.font.size = Pt(10.5)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    estilo.paragraph_format.space_after = Pt(6)
    estilo.paragraph_format.line_spacing = 1.15
    for nombre, tam in (("Heading 1", 14), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        h = doc.styles[nombre]
        h.font.name = FUENTE
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.color.rgb = COLOR_TITULO
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        h.paragraph_format.space_before = Pt(14 if nombre == "Heading 1" else 10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True


def campo(parrafo, instruccion):
    """Inserta un campo de Word (PAGE, NUMPAGES...)."""
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, instruccion), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)
    run.font.size = Pt(8)
    run.font.color.rgb = COLOR_SUAVE
    return run


def texto_con_formato(parrafo, texto, tam=None, color=None):
    """Añade texto interpretando **negrita** y *cursiva*."""
    hueco = r"\[[A-ZÁÉÍÓÚÑ_ ]{4,}[^\]]*\]"
    partes = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto)
    for p in partes:
        if not p:
            continue
        negrita = cursiva = False
        if p.startswith("**") and p.endswith("**"):
            p, negrita = p[2:-2], True
        elif p.startswith("*") and p.endswith("*") and len(p) > 2:
            p, cursiva = p[1:-1], True
        # los huecos pendientes van en su propio fragmento, resaltados en amarillo
        for trozo in re.split(f"({hueco})", p):
            if not trozo:
                continue
            run = parrafo.add_run(trozo)
            run.bold, run.italic = negrita or None, cursiva or None
            if tam:
                run.font.size = Pt(tam)
            if color:
                run.font.color.rgb = color
            if re.fullmatch(hueco, trozo):
                run.font.highlight_color = 7  # amarillo


def bordes_tabla(tabla):
    tbl = tabla._tbl
    tblPr = tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "B8BEC8")
        bordes.append(b)
    tblPr.append(bordes)


def sombrear(celda, color="E8ECF2"):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


# ---------------------------------------------------------------- bloques

def portada(doc, meta):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta.get("titulo", "CONTRATO DE PRESTACIÓN DE SERVICIOS"))
    r.bold, r.font.size, r.font.color.rgb = True, Pt(20), COLOR_TITULO

    if meta.get("proyecto"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(meta["proyecto"])
        r.font.size, r.font.color.rgb = Pt(13), COLOR_SUAVE

    doc.add_paragraph()
    for etiqueta, clave in (("Proveedor", "proveedor"), ("Cliente", "cliente")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{etiqueta}: ")
        r.bold = True
        p.add_run(meta.get(clave, "[DATO PENDIENTE]"))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lugar = meta.get("lugar", "")
    f = fecha_larga(meta.get("fecha", date.today().isoformat()))
    p.add_run(f"{lugar + ', ' if lugar else ''}{f}").font.color.rgb = COLOR_SUAVE

    if meta.get("revision_juridica", "no").lower() in ("si", "sí", "true", "yes"):
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("REVISIÓN JURÍDICA RECOMENDADA")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(12), COLOR_AVISO

    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(AVISO_BORRADOR)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9), COLOR_SUAVE


def indice(doc, titulos):
    h = doc.add_paragraph(style="Heading 1")
    h.add_run("Índice")
    for nivel, t in titulos:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0 if nivel == 1 else 0.6)
        r = p.add_run(t)
        r.bold = nivel == 1
        r.font.size = Pt(10 if nivel == 1 else 9.5)


def pie_de_pagina(seccion):
    pie = seccion.footer
    pie.is_linked_to_previous = False
    p = pie.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(AVISO_BORRADOR + "   ·   Página ")
    r.font.size, r.font.color.rgb = Pt(8), COLOR_SUAVE
    campo(p, "PAGE")
    r = p.add_run(" de ")
    r.font.size, r.font.color.rgb = Pt(8), COLOR_SUAVE
    campo(p, "NUMPAGES")


def tabla_markdown(doc, filas):
    filas = [f for f in filas if not re.match(r"^\|\s*:?-{2,}", f)]
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    ncol = max(len(c) for c in celdas)
    t = doc.add_table(rows=len(celdas), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(t)
    for i, fila in enumerate(celdas):
        for j in range(ncol):
            celda = t.cell(i, j)
            celda.paragraphs[0].paragraph_format.space_after = Pt(2)
            texto_con_formato(celda.paragraphs[0], fila[j] if j < len(fila) else "", tam=9.5)
            if i == 0:
                sombrear(celda)
                for r in celda.paragraphs[0].runs:
                    r.bold = True
    doc.add_paragraph()


def cuerpo(doc, texto):
    lineas = texto.splitlines()
    i = 0
    buffer = []

    def volcar():
        if buffer:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            texto_con_formato(p, " ".join(buffer))
            buffer.clear()

    while i < len(lineas):
        l = lineas[i].rstrip()
        s = l.strip()
        if not s:
            volcar()
        elif s.startswith("### "):
            volcar()
            doc.add_paragraph(s[4:], style="Heading 3")
        elif s.startswith("## "):
            volcar()
            doc.add_paragraph(s[3:], style="Heading 2")
        elif s.startswith("# "):
            volcar()
            titulo = s[2:]
            if titulo.upper().startswith("ANEXO"):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_paragraph(titulo, style="Heading 1")
        elif s.startswith(("- ", "* ")) and not s.startswith("**"):
            volcar()
            p = doc.add_paragraph(style="List Bullet")
            texto_con_formato(p, s[2:])
        elif s.startswith("|"):
            volcar()
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(lineas[i].strip())
                i += 1
            tabla_markdown(doc, filas)
            continue
        elif s in ("---", "***"):
            volcar()
        else:
            buffer.append(s)
        i += 1
    volcar()


def firmas(doc, meta):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    texto_con_formato(p, "Y en prueba de conformidad, las partes firman el presente contrato "
                         "en el lugar y la fecha indicados.")
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (rol, clave) in enumerate((("EL PROVEEDOR", "proveedor"), ("EL CLIENTE", "cliente"))):
        c = t.cell(0, j)
        par = c.paragraphs[0]
        par.paragraph_format.keep_with_next = True
        r = par.add_run(rol)
        r.bold = True
        for _ in range(3):
            c.add_paragraph()
        c.add_paragraph("______________________________")
        c.add_paragraph(meta.get(f"{clave}_firmante") or meta.get(clave, "[DATO PENDIENTE]"))
        cargo = meta.get(f"{clave}_cargo")
        if cargo:
            c.add_paragraph(cargo)
        if meta.get(f"{clave}_firmante") and meta.get(f"{clave}_firmante") != meta.get(clave):
            c.add_paragraph(f"En nombre de {meta.get(clave)}")
        c.add_paragraph("Fecha: ____ / ____ / ________")


# ---------------------------------------------------------------- principal

def generar(entrada, salida, pdf=True, forzar_indice=None):
    meta, texto = leer_entrada(entrada)
    os.makedirs(salida, exist_ok=True)

    nombre = "Contrato_{}_{}_{}_{}".format(
        slug(meta.get("proveedor")), slug(meta.get("cliente")),
        slug(meta.get("proyecto"), 25), meta.get("fecha", date.today().isoformat()))
    ruta_docx = os.path.join(salida, nombre + ".docx")

    doc = Document()
    fuente_base(doc)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4
    for lado in ("top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(2.5))
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(2.5))
    sec.different_first_page_header_footer = True  # portada sin pie
    pie_de_pagina(sec)

    portada(doc, meta)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    titulos = []
    for l in texto.splitlines():
        if l.startswith("# "):
            titulos.append((1, l[2:].strip()))
        elif l.startswith("## "):
            titulos.append((2, l[3:].strip()))
    usar_indice = forzar_indice if forzar_indice is not None else len(titulos) > 14
    if usar_indice:
        indice(doc, titulos)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # título y lugar/fecha al inicio del cuerpo
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta.get("titulo", "CONTRATO DE PRESTACIÓN DE SERVICIOS"))
    r.bold, r.font.size, r.font.color.rgb = True, Pt(13), COLOR_TITULO
    lugar = meta.get("lugar", "[CIUDAD]")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"En {lugar}, a {fecha_larga(meta.get('fecha', date.today().isoformat()))}.")

    # el cuerpo se divide en: contrato principal | anexos (firmas antes de los anexos)
    m = re.search(r"^# +ANEXO", texto, re.M | re.I)
    principal, anexos = (texto[:m.start()], texto[m.start():]) if m else (texto, "")
    cuerpo(doc, principal)
    firmas(doc, meta)
    if anexos:
        cuerpo(doc, anexos)

    doc.save(ruta_docx)
    print(f"DOCX: {ruta_docx}")

    if pdf:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            print("AVISO: LibreOffice no disponible; no se ha generado el PDF.")
            return ruta_docx, None
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", salida,
                            ruta_docx], check=True, capture_output=True, timeout=180)
            ruta_pdf = os.path.splitext(ruta_docx)[0] + ".pdf"
            print(f"PDF:  {ruta_pdf}")
            return ruta_docx, ruta_pdf
        except Exception as e:
            print(f"AVISO: no se pudo convertir a PDF ({e}). Entrega el Word.")
    return ruta_docx, None


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Genera contrato DOCX + PDF desde Markdown.")
    ap.add_argument("entrada")
    ap.add_argument("--salida", default=".")
    ap.add_argument("--sin-pdf", action="store_true")
    ap.add_argument("--indice", choices=["si", "no"], default=None)
    a = ap.parse_args()
    if not os.path.exists(a.entrada):
        sys.exit(f"No existe el archivo: {a.entrada}")
    generar(a.entrada, a.salida, pdf=not a.sin_pdf,
            forzar_indice=None if a.indice is None else a.indice == "si")
