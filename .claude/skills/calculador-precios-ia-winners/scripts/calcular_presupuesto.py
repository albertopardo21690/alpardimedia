#!/usr/bin/env python3
"""
Calcula los paquetes Básico / Estándar / Premium de un presupuesto de servicios de IA
y, opcionalmente, crea una hoja de cálculo Excel con fórmulas vivas.

Aplica exactamente las fórmulas de references/metodo-calculo.md.

FORMATO DE datos.json
---------------------
{
  "cliente": "Clínica Dental Sonrisa",
  "proyecto": "Agente de WhatsApp para reservas",
  "tarifa": 45,                     # €/h que cobra la persona
  "tarifa_minima": {                # opcional: número (€/h) o datos para calcularla
    "neto_deseado": 1800, "cuota_autonomos": 300, "gastos_negocio": 150,
    "irpf": 0.15, "horas_trabajo_mes": 120, "pct_facturable": 0.6, "meses_trabajo": 11
  },
  "riesgo": 0.20,                   # contingencia (0.10 / 0.20 / 0.30)
  "urgencia": 0.0,                  # 0 / 0.25 / 0.50
  "margen": 0.20,                   # margen de beneficio (0.15 - 0.30)
  "recargo_gestion": 0.15,          # sobre costes de herramientas que gestiona el proveedor
  "valor": {"horas_ahorradas_mes": 20, "coste_hora_empleado": 18},   # opcional
  "recomendado": "Estándar",
  "paquetes": {
    "Básico": {
      "horas": {"Análisis y descubrimiento": 3, "Reuniones y gestión": 2, "Diseño": 2,
                "Desarrollo e integración": 10, "Pruebas": 4,
                "Documentación y formación": 2, "Entrega y puesta en marcha": 1},
      "costes_unicos": [{"concepto": "Número WhatsApp", "importe": 0, "gestiona": "cliente"}],
      "costes_recurrentes": [{"concepto": "API IA (supuesto 500 conv.)", "importe_mes": 15,
                              "gestiona": "proveedor"}],
      "horas_mantenimiento_mes": 1,
      "plazo": "2 semanas",
      "riesgo": 0.10                # opcional: sustituye al riesgo general para este paquete
    },
    "Estándar": { ... },
    "Premium": { ... }
  }
}

"gestiona": "proveedor" → el coste se incluye en el precio con el recargo de gestión.
"gestiona": "cliente"   → el cliente lo paga directamente; se lista pero no se suma.

USO
---
    python3 calcular_presupuesto.py datos.json
    python3 calcular_presupuesto.py datos.json --excel Presupuesto.xlsx
    python3 calcular_presupuesto.py datos.json --json   # salida en JSON
"""

import argparse
import json
import math
import re
import sys


def redondear(x, paso):
    return int(math.ceil(round(x, 6) / paso) * paso) if x > 0 else 0


def eur(x):
    s = f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return (s[:-3] if s.endswith(",00") else s) + " €"


def pct(x):
    return f"{x * 100:.0f} %"


def tarifa_minima(dato):
    if dato is None:
        return None, None
    if isinstance(dato, (int, float)):
        return float(dato), None
    d = {"cuota_autonomos": 300, "gastos_negocio": 150, "irpf": 0.15,
         "horas_trabajo_mes": 120, "pct_facturable": 0.6, "meses_trabajo": 11, **dato}
    if "neto_deseado" not in d:
        return None, None
    bruto = d["neto_deseado"] / (1 - d["irpf"]) + d["cuota_autonomos"] + d["gastos_negocio"]
    ajustado = bruto * 12 / d["meses_trabajo"]
    horas = d["horas_trabajo_mes"] * d["pct_facturable"]
    return ajustado / horas, {"bruto_necesario_mes": bruto, "bruto_ajustado": ajustado,
                              "horas_facturables": horas, **d}


def calcular(datos):
    tarifa = float(datos["tarifa"])
    tmin, detalle_tmin = tarifa_minima(datos.get("tarifa_minima"))
    margen = datos.get("margen", 0.20)
    urg = datos.get("urgencia", 0.0)
    gest = datos.get("recargo_gestion", 0.15)
    res = {"tarifa": tarifa, "tarifa_minima": tmin, "detalle_tarifa_minima": detalle_tmin,
           "avisos": [], "paquetes": {}}

    if tmin and tarifa < tmin:
        res["avisos"].append(
            f"La tarifa de {eur(tarifa)}/h está por debajo de la mínima sostenible "
            f"({eur(tmin)}/h). Con esta tarifa el negocio pierde dinero.")
    if margen < 0.15:
        res["avisos"].append(f"Margen del {pct(margen)}: por debajo del 15 % recomendado.")

    for nombre, p in datos["paquetes"].items():
        riesgo = p.get("riesgo", datos.get("riesgo", 0.20))
        horas = p.get("horas", {})
        h_total = sum(horas.values())
        coste_trabajo = h_total * tarifa
        con_riesgo = coste_trabajo * (1 + riesgo)
        precio_trabajo = con_riesgo * (1 + urg) * (1 + margen)

        cu = p.get("costes_unicos", [])
        cu_prov = sum(c["importe"] for c in cu if c.get("gestiona", "cliente") == "proveedor")
        cu_cli = sum(c["importe"] for c in cu if c.get("gestiona", "cliente") != "proveedor")
        impl_bruto = precio_trabajo + cu_prov * (1 + gest)
        impl = redondear(impl_bruto, 50)

        cr = p.get("costes_recurrentes", [])
        cr_prov = sum(c["importe_mes"] for c in cr if c.get("gestiona", "cliente") == "proveedor")
        cr_cli = sum(c["importe_mes"] for c in cr if c.get("gestiona", "cliente") != "proveedor")
        h_mant = p.get("horas_mantenimiento_mes", 0)
        mens_bruto = h_mant * tarifa * (1 + margen) + cr_prov * (1 + gest)
        mens = redondear(mens_bruto, 10)

        t_suelo = tmin if tmin else tarifa
        minimo = h_total * t_suelo * (1 + riesgo) + cu_prov
        rango = (max(minimo, impl * 0.90), impl * 1.10)

        # fases: comprobar que no falten pruebas / análisis / reuniones
        claves = " ".join(horas.keys()).lower()
        faltan = [f for f, pat in (("análisis", r"an[aá]lisis|descubr"), ("reuniones", r"reuni|gesti"),
                                   ("pruebas", r"prueba|test"))
                  if not re.search(pat, claves)]
        if faltan:
            res["avisos"].append(f"{nombre}: no hay horas de {', '.join(faltan)}. Revisa si faltan.")
        if h_total and horas:
            dev = sum(v for k, v in horas.items() if re.search(r"desarroll|integra", k.lower()))
            if dev / h_total > 0.65:
                res["avisos"].append(
                    f"{nombre}: el desarrollo es el {pct(dev / h_total)} del total. "
                    "Suele ser el 40-55 %; puede que falten horas de análisis, pruebas o reuniones.")

        r = {"horas": horas, "horas_totales": h_total, "riesgo": riesgo,
             "coste_trabajo": coste_trabajo, "con_riesgo": con_riesgo,
             "precio_trabajo": precio_trabajo,
             "costes_unicos_proveedor": cu_prov, "costes_unicos_cliente": cu_cli,
             "implementacion_sin_redondear": impl_bruto, "implementacion": impl,
             "costes_recurrentes_proveedor": cr_prov, "costes_recurrentes_cliente": cr_cli,
             "horas_mantenimiento_mes": h_mant, "mensualidad_sin_redondear": mens_bruto,
             "mensualidad": mens, "minimo_sostenible": minimo,
             "rango_negociacion": rango, "plazo": p.get("plazo", ""),
             "costes_unicos": cu, "costes_recurrentes": cr}

        v = datos.get("valor")
        if v and v.get("horas_ahorradas_mes") and v.get("coste_hora_empleado"):
            ahorro = v["horas_ahorradas_mes"] * v["coste_hora_empleado"]
            neto = ahorro - mens - cr_cli
            r["ahorro_mensual_estimado"] = ahorro
            r["meses_para_recuperar"] = (impl + cu_cli) / neto if neto > 0 else None
        res["paquetes"][nombre] = r
    return res


def informe(datos, res):
    out = [f"# Presupuesto: {datos.get('proyecto', '')} · {datos.get('cliente', '')}", ""]
    out.append(f"Tarifa: {eur(res['tarifa'])}/h · Riesgo general: {pct(datos.get('riesgo', 0.2))} · "
               f"Urgencia: {pct(datos.get('urgencia', 0))} · Margen: {pct(datos.get('margen', 0.2))} · "
               f"Recargo gestión herramientas: {pct(datos.get('recargo_gestion', 0.15))}")
    if res["tarifa_minima"]:
        out.append(f"Tarifa mínima sostenible: {eur(res['tarifa_minima'])}/h")
        d = res["detalle_tarifa_minima"]
        if d:
            out.append(f"  ({eur(d['neto_deseado'])} / (1 - {pct(d['irpf'])}) + {eur(d['cuota_autonomos'])} "
                       f"+ {eur(d['gastos_negocio'])} = {eur(d['bruto_necesario_mes'])} × 12/{d['meses_trabajo']} "
                       f"= {eur(d['bruto_ajustado'])} / {d['horas_facturables']:.0f} h facturables)")
    out.append("")
    for n, r in res["paquetes"].items():
        marca = " ⭐ recomendado" if n == datos.get("recomendado") else ""
        out.append(f"## {n}{marca}")
        for fase, h in r["horas"].items():
            out.append(f"- {fase}: {h:g} h")
        out.append(f"- **Total: {r['horas_totales']:g} h** × {eur(res['tarifa'])} = {eur(r['coste_trabajo'])}")
        out.append(f"- + riesgo {pct(r['riesgo'])} = {eur(r['con_riesgo'])}"
                   f" → × urgencia y margen = {eur(r['precio_trabajo'])}")
        if r["costes_unicos_proveedor"]:
            out.append(f"- + costes únicos gestionados {eur(r['costes_unicos_proveedor'])} (+ recargo)")
        out.append(f"- **Implementación: {eur(r['implementacion'])} + IVA**")
        if r["mensualidad"]:
            out.append(f"- **Mensualidad: {eur(r['mensualidad'])} + IVA** "
                       f"({r['horas_mantenimiento_mes']:g} h/mes + herramientas gestionadas "
                       f"{eur(r['costes_recurrentes_proveedor'])})")
        if r["costes_recurrentes_cliente"] or r["costes_unicos_cliente"]:
            out.append(f"- Herramientas que paga el cliente directamente: "
                       f"{eur(r['costes_unicos_cliente'])} únicos + {eur(r['costes_recurrentes_cliente'])}/mes")
        out.append(f"- Mínimo sostenible: {eur(r['minimo_sostenible'])} · Rango de negociación: "
                   f"{eur(r['rango_negociacion'][0])} – {eur(r['rango_negociacion'][1])}")
        if r.get("ahorro_mensual_estimado"):
            m = r["meses_para_recuperar"]
            out.append(f"- Valor estimado (supuesto): ahorro de {eur(r['ahorro_mensual_estimado'])}/mes → "
                       + (f"recupera la inversión en ~{m:.1f} meses".replace(".", ",") if m else
                          "el ahorro no cubre la mensualidad: revisa el alcance"))
        if r["plazo"]:
            out.append(f"- Plazo estimado: {r['plazo']}")
        out.append("")
    if res["avisos"]:
        out.append("## Avisos")
        out += [f"- ⚠️ {a}" for a in res["avisos"]]
    return "\n".join(out)


def excel(datos, res, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    azul = PatternFill("solid", fgColor="1F2A44")
    suave = PatternFill("solid", fgColor="E8ECF2")
    editable = PatternFill("solid", fgColor="FFF4CC")
    blanco = Font(color="FFFFFF", bold=True)
    negrita = Font(bold=True)
    fino = Side(style="thin", color="B8BEC8")
    borde = Border(top=fino, bottom=fino, left=fino, right=fino)
    EUR = '#,##0.00 "€"'
    PCT = "0%"

    wb = Workbook()
    ps = wb.active
    ps.title = "Parámetros"
    ps["A1"] = f"Presupuesto · {datos.get('proyecto', '')} · {datos.get('cliente', '')}"
    ps["A1"].font = Font(bold=True, size=14, color="1F2A44")
    ps["A2"] = "Las celdas amarillas se pueden editar: el resto se recalcula solo."
    ps["A2"].font = Font(italic=True, color="595F6B")
    params = [("Tarifa (€/h)", res["tarifa"], EUR),
              ("Tarifa mínima sostenible (€/h)", res["tarifa_minima"] or res["tarifa"], EUR),
              ("Urgencia", datos.get("urgencia", 0.0), PCT),
              ("Margen de beneficio", datos.get("margen", 0.20), PCT),
              ("Recargo gestión herramientas", datos.get("recargo_gestion", 0.15), PCT)]
    for i, (k, v, fmt) in enumerate(params, start=4):
        ps.cell(i, 1, k).font = negrita
        c = ps.cell(i, 2, v)
        c.number_format, c.fill, c.border = fmt, editable, borde
    ps.column_dimensions["A"].width = 34
    ps.column_dimensions["B"].width = 16
    P = {"tarifa": "Parámetros!$B$4", "tmin": "Parámetros!$B$5", "urg": "Parámetros!$B$6",
         "margen": "Parámetros!$B$7", "gest": "Parámetros!$B$8"}

    resumen = []
    for nombre, r in res["paquetes"].items():
        ws = wb.create_sheet(nombre[:31])
        ws["A1"] = f"Paquete {nombre}"
        ws["A1"].font = Font(bold=True, size=13, color="1F2A44")
        ws["A3"], ws["B3"], ws["C3"] = "Fase", "Horas", "Coste"
        for c in ("A3", "B3", "C3"):
            ws[c].fill, ws[c].font = azul, blanco
        fila = 4
        for fase, h in r["horas"].items():
            ws.cell(fila, 1, fase).border = borde
            c = ws.cell(fila, 2, h)
            c.fill, c.border = editable, borde
            c = ws.cell(fila, 3, f"=B{fila}*{P['tarifa']}")
            c.number_format, c.border = EUR, borde
            fila += 1
        f_ini, f_fin = 4, fila - 1
        ws.cell(fila, 1, "Total").font = negrita
        ws.cell(fila, 2, f"=SUM(B{f_ini}:B{f_fin})").font = negrita
        c = ws.cell(fila, 3, f"=SUM(C{f_ini}:C{f_fin})")
        c.number_format, c.font = EUR, negrita
        f_tot = fila
        fila += 2

        def linea(etq, formula, fmt=EUR, bold=False, edit=False):
            nonlocal fila
            ws.cell(fila, 1, etq).font = Font(bold=bold)
            c = ws.cell(fila, 3, formula)
            c.number_format = fmt
            c.border = borde
            if bold:
                c.font = negrita
            if edit:
                c.fill = editable
            fila += 1
            return f"C{fila - 1}"

        riesgo = linea("Contingencia por riesgo", r["riesgo"], PCT, edit=True)
        con_r = linea("Trabajo + riesgo", f"=C{f_tot}*(1+{riesgo})")
        p_trab = linea("× urgencia × margen", f"=C{con_r[1:]}*(1+{P['urg']})*(1+{P['margen']})")
        fila += 1

        # costes
        ws.cell(fila, 1, "Costes únicos").fill = suave
        ws.cell(fila, 2, "Gestiona").fill = suave
        ws.cell(fila, 3, "Importe").fill = suave
        fila += 1
        cu_ini = fila
        for c_ in r["costes_unicos"] or [{"concepto": "(ninguno)", "importe": 0, "gestiona": "cliente"}]:
            ws.cell(fila, 1, c_["concepto"])
            ws.cell(fila, 2, c_.get("gestiona", "cliente")).fill = editable
            x = ws.cell(fila, 3, c_["importe"])
            x.number_format, x.fill = EUR, editable
            fila += 1
        cu_fin = fila - 1
        cu_prov = linea("Costes únicos gestionados por el proveedor",
                        f'=SUMIF(B{cu_ini}:B{cu_fin},"proveedor",C{cu_ini}:C{cu_fin})')
        fila += 1
        ws.cell(fila, 1, "Costes recurrentes (mes)").fill = suave
        ws.cell(fila, 2, "Gestiona").fill = suave
        ws.cell(fila, 3, "€/mes").fill = suave
        fila += 1
        cr_ini = fila
        for c_ in r["costes_recurrentes"] or [{"concepto": "(ninguno)", "importe_mes": 0, "gestiona": "cliente"}]:
            ws.cell(fila, 1, c_["concepto"])
            ws.cell(fila, 2, c_.get("gestiona", "cliente")).fill = editable
            x = ws.cell(fila, 3, c_["importe_mes"])
            x.number_format, x.fill = EUR, editable
            fila += 1
        cr_fin = fila - 1
        cr_prov = linea("Recurrentes gestionados por el proveedor",
                        f'=SUMIF(B{cr_ini}:B{cr_fin},"proveedor",C{cr_ini}:C{cr_fin})')
        cr_cli = linea("Recurrentes que paga el cliente directamente",
                       f'=SUMIF(B{cr_ini}:B{cr_fin},"<>proveedor",C{cr_ini}:C{cr_fin})')
        fila += 1
        impl = linea("IMPLEMENTACIÓN (+ IVA)",
                     f"=CEILING({p_trab}+{cu_prov}*(1+{P['gest']}),50)", bold=True)
        h_mant = linea("Horas de mantenimiento al mes", r["horas_mantenimiento_mes"], "0.0", edit=True)
        mens = linea("MENSUALIDAD (+ IVA)",
                     f"=IF({h_mant}*{P['tarifa']}*(1+{P['margen']})+{cr_prov}*(1+{P['gest']})>0,"
                     f"CEILING({h_mant}*{P['tarifa']}*(1+{P['margen']})+{cr_prov}*(1+{P['gest']}),10),0)",
                     bold=True)
        minimo = linea("Precio mínimo sostenible", f"=C{f_tot}/{P['tarifa']}*{P['tmin']}*(1+{riesgo})+{cu_prov}")
        r_min = linea("Rango de negociación: desde", f"=MAX({minimo},{impl}*0.9)")
        r_max = linea("Rango de negociación: hasta", f"={impl}*1.1")
        ws.cell(fila + 1, 1, f"Plazo estimado: {r['plazo']}").font = Font(italic=True)
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 16
        resumen.append((nombre, f"'{ws.title}'!B{f_tot}", f"'{ws.title}'!{impl}",
                        f"'{ws.title}'!{mens}", f"'{ws.title}'!{cr_cli}",
                        f"'{ws.title}'!{minimo}", f"'{ws.title}'!{r_min}", f"'{ws.title}'!{r_max}",
                        r["plazo"]))

    rs = wb.create_sheet("Resumen", 1)
    cab = ["Paquete", "Horas", "Implementación", "Mensualidad", "Herramientas cliente/mes",
           "Mínimo sostenible", "Negociación desde", "Negociación hasta", "Plazo"]
    for j, t in enumerate(cab, start=1):
        c = rs.cell(1, j, t)
        c.fill, c.font = azul, blanco
        c.alignment = Alignment(wrap_text=True, vertical="center")
        rs.column_dimensions[get_column_letter(j)].width = 16
    rs.column_dimensions["A"].width = 18
    for i, fila_r in enumerate(resumen, start=2):
        nombre = fila_r[0] + (" ⭐" if fila_r[0] == datos.get("recomendado") else "")
        rs.cell(i, 1, nombre).font = negrita
        for j, ref in enumerate(fila_r[1:8], start=2):
            c = rs.cell(i, j, f"={ref}")
            c.number_format = "0" if j == 2 else EUR
            c.border = borde
        rs.cell(i, 9, fila_r[8])
    rs.cell(len(resumen) + 3, 1, "Precios sin IVA. Los costes de herramientas son estimaciones: "
                                 "confirma los precios actuales.").font = Font(italic=True, color="595F6B")
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("datos")
    ap.add_argument("--excel")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    with open(a.datos, encoding="utf-8") as f:
        datos = json.load(f)
    if "tarifa" not in datos or "paquetes" not in datos:
        sys.exit("Faltan 'tarifa' o 'paquetes' en el archivo de datos.")
    res = calcular(datos)
    if a.json:
        print(json.dumps(res, ensure_ascii=False, indent=2, default=str))
    else:
        print(informe(datos, res))
    if a.excel:
        print(f"\nExcel: {excel(datos, res, a.excel)}")
